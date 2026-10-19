import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import win32com.client

# Paths
base_dir = r"d:\ITR Project\lamiwood-ecom-website\report"
assets_dir = os.path.join(base_dir, "assets")
shots_dir = os.path.join(base_dir, "screenshots")
sample_extracted = os.path.join(base_dir, "sample_ppt_extracted")
output_pptx = os.path.join(base_dir, "Lamiwood_Industrial_Training_Presentation.pptx")
output_pdf = os.path.join(base_dir, "Lamiwood_Industrial_Training_Presentation.pdf")
slides_img_dir = os.path.join(base_dir, "presentation_slides")
os.makedirs(slides_img_dir, exist_ok=True)

# Image paths
skit_logo = os.path.join(assets_dir, "skit_logo_crisp.png")
if not os.path.exists(skit_logo):
    skit_logo = os.path.join(sample_extracted, "page_1_img_1.jpeg")

swami_portrait = os.path.join(assets_dir, "swami_portrait_crisp.png")
if not os.path.exists(swami_portrait):
    swami_portrait = os.path.join(sample_extracted, "page_1_img_2.jpeg")

cert_img = os.path.join(assets_dir, "aditya_swami_certificate.png")
tech_stack_img = os.path.join(assets_dir, "tech_stack_lamiwood.png")

shot_home = os.path.join(shots_dir, "home.png")
shot_products = os.path.join(shots_dir, "products.png")
shot_detail = os.path.join(shots_dir, "product_detail.png")
shot_pattern = os.path.join(shots_dir, "pattern_designer.png")
shot_calc = os.path.join(shots_dir, "calculator.png")
shot_compare = os.path.join(shots_dir, "compare.png")

# Colors
RED_THEME = RGBColor(192, 0, 0)       # #C00000
DARK_RED = RGBColor(144, 0, 0)        # #900000
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
BODY_TEXT = RGBColor(20, 20, 20)
BORDER_GRAY = RGBColor(210, 210, 210)
CARD_BG = RGBColor(248, 249, 250)

def create_9_slide_presentation():
    prs = Presentation()
    # 16:9 widescreen: 13.333 inches x 7.5 inches
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_skit_logo(slide, left=Inches(0.38), top=Inches(0.18), width=Inches(1.25), height=Inches(1.25)):
        if os.path.exists(skit_logo):
            slide.shapes.add_picture(skit_logo, left, top, width, height)

    def add_header_banner(slide, title_text):
        banner = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(2.473), Inches(0.24), Inches(10.3), Inches(0.723)
        )
        banner.fill.solid()
        banner.fill.fore_color.rgb = RED_THEME
        banner.line.fill.background()
        
        tf = banner.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title_text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Times New Roman"
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = WHITE

    def add_footer_banner(slide):
        footer = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.0), Inches(7.233), Inches(13.333333), Inches(0.267)
        )
        footer.fill.solid()
        footer.fill.fore_color.rgb = RED_THEME
        footer.line.fill.background()

        # Left text
        tb_left = slide.shapes.add_textbox(Inches(0.25), Inches(7.21), Inches(4.0), Inches(0.28))
        tf_l = tb_left.text_frame
        tf_l.word_wrap = False
        tf_l.margin_top = Inches(0)
        tf_l.margin_bottom = Inches(0)
        p_l = tf_l.paragraphs[0]
        p_l.text = "SKIT, JAIPUR"
        p_l.font.name = "Calibri"
        p_l.font.size = Pt(11.5)
        p_l.font.bold = True
        p_l.font.color.rgb = WHITE

        # Right text
        tb_right = slide.shapes.add_textbox(Inches(9.08), Inches(7.21), Inches(4.0), Inches(0.28))
        tf_r = tb_right.text_frame
        tf_r.word_wrap = False
        tf_r.margin_top = Inches(0)
        tf_r.margin_bottom = Inches(0)
        p_r = tf_r.paragraphs[0]
        p_r.text = "www.skit.ac.in"
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.font.name = "Calibri"
        p_r.font.size = Pt(11.5)
        p_r.font.bold = True
        p_r.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s1, Inches(0.38), Inches(0.19), Inches(1.25), Inches(1.25))
    if os.path.exists(swami_portrait):
        s1.shapes.add_picture(swami_portrait, Inches(12.0), Inches(0.19), Inches(0.85), Inches(1.35))

    tb_c = s1.shapes.add_textbox(Inches(1.8), Inches(0.35), Inches(9.8), Inches(1.2))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p1 = tf_c.paragraphs[0]
    p1.text = "Swami Keshvanand Institute of Technology, Management & Gramothan, Jaipur"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Times New Roman"
    p1.font.size = Pt(19)
    p1.font.bold = True
    p1.font.color.rgb = DARK_RED

    p2 = tf_c.add_paragraph()
    p2.text = "(An Autonomous Institute Affiliated to Rajasthan Technical University, Kota)"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Times New Roman"
    p2.font.size = Pt(16.5)
    p2.font.bold = True
    p2.font.color.rgb = DARK_RED
    p2.space_before = Pt(4)

    tb_t = s1.shapes.add_textbox(Inches(1.5), Inches(1.65), Inches(10.33), Inches(1.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p_t1 = tf_t.paragraphs[0]
    p_t1.text = "Industrial Training"
    p_t1.alignment = PP_ALIGN.CENTER
    p_t1.font.name = "Times New Roman"
    p_t1.font.size = Pt(30)
    p_t1.font.bold = True
    p_t1.font.color.rgb = BLACK

    p_t2 = tf_t.add_paragraph()
    p_t2.text = "Presentation"
    p_t2.alignment = PP_ALIGN.CENTER
    p_t2.font.name = "Times New Roman"
    p_t2.font.size = Pt(30)
    p_t2.font.bold = True
    p_t2.font.color.rgb = BLACK
    p_t2.space_before = Pt(3)

    tb_top = s1.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(11.33), Inches(1.6))
    tf_top = tb_top.text_frame
    tf_top.word_wrap = True
    p_top1 = tf_top.paragraphs[0]
    p_top1.text = "TOPIC : LAMIWOOD – MODERN E-COMMERCE PLATFORM"
    p_top1.alignment = PP_ALIGN.CENTER
    p_top1.font.name = "Arial Black"
    p_top1.font.size = Pt(25)
    p_top1.font.bold = True
    p_top1.font.color.rgb = BLACK

    p_top2 = tf_top.add_paragraph()
    p_top2.text = "FOR LUXURY ARCHITECTURAL SURFACES"
    p_top2.alignment = PP_ALIGN.CENTER
    p_top2.font.name = "Arial Black"
    p_top2.font.size = Pt(23)
    p_top2.font.bold = True
    p_top2.font.color.rgb = BLACK
    p_top2.space_before = Pt(4)

    # Submitted To
    tb_sto = s1.shapes.add_textbox(Inches(0.55), Inches(5.2), Inches(4.8), Inches(1.7))
    tf_sto = tb_sto.text_frame
    tf_sto.word_wrap = True
    p_sto1 = tf_sto.paragraphs[0]
    p_sto1.text = "Submitted To:"
    p_sto1.font.name = "Times New Roman"
    p_sto1.font.size = Pt(20)
    p_sto1.font.bold = True
    p_sto1.font.color.rgb = BLACK

    p_sto2 = tf_sto.add_paragraph()
    p_sto2.text = "Name:- Mr Vikash Dhakar"
    p_sto2.font.name = "Times New Roman"
    p_sto2.font.size = Pt(19)
    p_sto2.font.bold = True
    p_sto2.font.color.rgb = BLACK
    p_sto2.space_before = Pt(6)

    p_sto3 = tf_sto.add_paragraph()
    p_sto3.text = "Assistant Professor"
    p_sto3.font.name = "Times New Roman"
    p_sto3.font.size = Pt(19)
    p_sto3.font.bold = True
    p_sto3.font.color.rgb = BLACK
    p_sto3.space_before = Pt(4)

    # Submitted By
    tb_sby = s1.shapes.add_textbox(Inches(8.8), Inches(5.2), Inches(4.2), Inches(1.7))
    tf_sby = tb_sby.text_frame
    tf_sby.word_wrap = True
    p_sby1 = tf_sby.paragraphs[0]
    p_sby1.text = "Submitted By:"
    p_sby1.font.name = "Times New Roman"
    p_sby1.font.size = Pt(20)
    p_sby1.font.bold = True
    p_sby1.font.color.rgb = BLACK

    p_sby2 = tf_sby.add_paragraph()
    p_sby2.text = "Name:- Aditya Swami"
    p_sby2.font.name = "Times New Roman"
    p_sby2.font.size = Pt(19)
    p_sby2.font.bold = True
    p_sby2.font.color.rgb = BLACK
    p_sby2.space_before = Pt(6)

    p_sby3 = tf_sby.add_paragraph()
    p_sby3.text = "Roll No.:- 24ESKCS024"
    p_sby3.font.name = "Times New Roman"
    p_sby3.font.size = Pt(19)
    p_sby3.font.bold = True
    p_sby3.font.color.rgb = BLACK
    p_sby3.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: INDEX
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s2)
    add_header_banner(s2, "Index")
    add_footer_banner(s2)

    tb_idx = s2.shapes.add_textbox(Inches(2.7), Inches(1.5), Inches(8.5), Inches(5.3))
    tf_idx = tb_idx.text_frame
    tf_idx.word_wrap = True
    index_items = [
        "1. Company Introduction & Certificate",
        "2. Technology Stack",
        "3. Goals of Training & Project Overview",
        "4. Working Project: Storefront & Catalog",
        "5. Working Project: Interactive Design Tools",
        "6. Future Scope & Conclusion",
        "7. Q&A"
    ]
    for i, item in enumerate(index_items):
        p = tf_idx.paragraphs[0] if i == 0 else tf_idx.add_paragraph()
        p.text = item
        p.font.name = "Times New Roman"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = BODY_TEXT
        if i > 0:
            p.space_before = Pt(12)

    # =========================================================================
    # SLIDE 3: COMPANY INTRODUCTION & CERTIFICATE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s3)
    add_header_banner(s3, "Company Introduction & Certificate")
    add_footer_banner(s3)

    # Left text box
    tb_c_info = s3.shapes.add_textbox(Inches(1.8), Inches(1.3), Inches(5.6), Inches(5.6))
    tf_ci = tb_c_info.text_frame
    tf_ci.word_wrap = True

    p = tf_ci.paragraphs[0]
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.line_spacing = 1.15
    r = p.add_run()
    r.text = "Linkstatic Techno "
    r.font.bold = True
    r = p.add_run()
    r.text = "is a premier technology consulting and digital product engineering organization specializing in modern full-stack web applications, reactive user interfaces, and bespoke enterprise solutions."

    p = tf_ci.add_paragraph()
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.line_spacing = 1.15
    p.space_before = Pt(10)
    r = p.add_run()
    r.text = "During my 6-weeks industrial training (15 May 2026 – 29 June 2026), I worked as a "
    r = p.add_run()
    r.text = "Full Stack Web & Frontend Engineering Intern"
    r.font.bold = True
    r = p.add_run()
    r.text = ", under the mentorship of "
    r = p.add_run()
    r.text = "Mr. Aditya Pareek "
    r.font.bold = True
    r = p.add_run()
    r.text = "(Founder & CEO)."

    p = tf_ci.add_paragraph()
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.line_spacing = 1.15
    p.space_before = Pt(10)
    r = p.add_run()
    r.text = "Primary Focus: "
    r.font.bold = True
    r = p.add_run()
    r.text = "Architecting the "
    r = p.add_run()
    r.text = "Lamiwood "
    r.font.bold = True
    r = p.add_run()
    r.text = "e-commerce web platform using "
    r = p.add_run()
    r.text = "React 19, TypeScript, and Tailwind CSS v4"
    r.font.bold = True
    r = p.add_run()
    r.text = ", implementing interactive SVG layout engines, dynamic calculators, and state persistence."

    # Right Certificate Image
    if os.path.exists(cert_img):
        pic = s3.shapes.add_picture(
            cert_img,
            Inches(7.6), Inches(1.4), Inches(5.35), Inches(5.2)
        )
        pic.line.color.rgb = BORDER_GRAY
        pic.line.width = Pt(1)

    # =========================================================================
    # SLIDE 4: TECHNOLOGY STACK
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s4)
    add_footer_banner(s4)

    if os.path.exists(tech_stack_img):
        tech_pic = s4.shapes.add_picture(
            tech_stack_img,
            Inches(1.8), Inches(0.4), Inches(10.2), Inches(6.4)
        )
        tech_pic.line.color.rgb = BORDER_GRAY
        tech_pic.line.width = Pt(1)

    # =========================================================================
    # SLIDE 5: GOALS OF TRAINING & PROJECT OVERVIEW
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s5)
    add_header_banner(s5, "Training Goals & Project Overview")
    add_footer_banner(s5)

    # Left: Goals
    tb_g = s5.shapes.add_textbox(Inches(1.8), Inches(1.3), Inches(5.3), Inches(5.6))
    tf_g = tb_g.text_frame
    tf_g.word_wrap = True

    p = tf_g.paragraphs[0]
    p.text = "Goal of Training"
    p.font.name = "Times New Roman"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = RED_THEME

    goals_short = [
        ("1. Real-World Experience: ", "Engineering a production-ready, full-stack architectural surfaces e-commerce application."),
        ("2. Modern Frontend Mastery: ", "Applying React 19, TypeScript, and Tailwind CSS v4 for luxury responsive UI systems."),
        ("3. Interactive Tool Engineering: ", "Developing live SVG pattern simulators, wastage calculators, and comparison matrices."),
        ("4. Professional Practices: ", "Mastering Git version control, GitHub collaboration, Context API state, and performance profiling.")
    ]
    for num, txt in goals_short:
        p = tf_g.add_paragraph()
        p.font.name = "Calibri"
        p.font.size = Pt(15.5)
        p.line_spacing = 1.15
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = num
        r1.font.bold = True
        r2 = p.add_run()
        r2.text = txt

    # Right: Overview
    tb_o = s5.shapes.add_textbox(Inches(7.4), Inches(1.3), Inches(5.5), Inches(5.6))
    tf_o = tb_o.text_frame
    tf_o.word_wrap = True

    p = tf_o.paragraphs[0]
    p.text = "Lamiwood Platform Overview"
    p.font.name = "Times New Roman"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = RED_THEME

    p = tf_o.add_paragraph()
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.line_spacing = 1.18
    p.space_before = Pt(8)
    r = p.add_run()
    r.text = "Lamiwood "
    r.font.bold = True
    r = p.add_run()
    r.text = "is an advanced digital showroom designed for "
    r = p.add_run()
    r.text = "architects, interior designers, and homeowners "
    r.font.bold = True
    r = p.add_run()
    r.text = "to discover and specify luxury interior surfaces—including engineered hardwood, acoustic wall panels, and outdoor decking."

    p = tf_o.add_paragraph()
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.line_spacing = 1.18
    p.space_before = Pt(8)
    r = p.add_run()
    r.text = "Key Problem Solved: "
    r.font.bold = True
    r = p.add_run()
    r.text = "Replaces fragmented offline sample distribution with instant digital exploration, dynamic pattern previews, material estimation, and a "
    r = p.add_run()
    r.text = "Free Physical Sample Box "
    r.font.bold = True
    r = p.add_run()
    r.text = "fulfillment workflow."

    p = tf_o.add_paragraph()
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.line_spacing = 1.18
    p.space_before = Pt(8)
    r = p.add_run()
    r.text = "Architecture: "
    r.font.bold = True
    r = p.add_run()
    r.text = "Component-driven Single Page Application with persistent client-side state (React Context API + LocalStorage), Framer Motion reveals, and sub-second page transitions."

    # =========================================================================
    # SLIDE 6: WORKING PROJECT – STOREFRONT & CATALOG
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s6)
    add_header_banner(s6, "Working Project: Storefront & Catalog")
    add_footer_banner(s6)

    # 2 Screenshots side-by-side with labels
    # Screenshot 1: Home Page
    if os.path.exists(shot_home):
        pic1 = s6.shapes.add_picture(
            shot_home,
            Inches(1.8), Inches(1.35), Inches(5.3), Inches(4.3)
        )
        pic1.line.color.rgb = BORDER_GRAY
        pic1.line.width = Pt(1)

    tb_lbl1 = s6.shapes.add_textbox(Inches(1.8), Inches(5.75), Inches(5.3), Inches(1.2))
    tf_l1 = tb_lbl1.text_frame
    tf_l1.word_wrap = True
    p = tf_l1.paragraphs[0]
    p.text = "Home Page & Brand Identity"
    p.font.name = "Times New Roman"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RED_THEME
    p = tf_l1.add_paragraph()
    p.text = "Luxury hero presentation, curated surface categories, and quick navigation."
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.color.rgb = BODY_TEXT

    # Screenshot 2: Product Catalog
    if os.path.exists(shot_products):
        pic2 = s6.shapes.add_picture(
            shot_products,
            Inches(7.5), Inches(1.35), Inches(5.3), Inches(4.3)
        )
        pic2.line.color.rgb = BORDER_GRAY
        pic2.line.width = Pt(1)

    tb_lbl2 = s6.shapes.add_textbox(Inches(7.5), Inches(5.75), Inches(5.3), Inches(1.2))
    tf_l2 = tb_lbl2.text_frame
    tf_l2.word_wrap = True
    p = tf_l2.paragraphs[0]
    p.text = "Product Catalog & Faceted Filters"
    p.font.name = "Times New Roman"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RED_THEME
    p = tf_l2.add_paragraph()
    p.text = "Multi-attribute category filters, price sliders, technical badges, and responsive grid."
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.color.rgb = BODY_TEXT

    # =========================================================================
    # SLIDE 7: WORKING PROJECT – INTERACTIVE DESIGN TOOLS
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s7)
    add_header_banner(s7, "Working Project: Interactive Tools")
    add_footer_banner(s7)

    # Screenshot 1: Pattern Designer
    if os.path.exists(shot_pattern):
        pic1 = s7.shapes.add_picture(
            shot_pattern,
            Inches(1.8), Inches(1.35), Inches(5.3), Inches(4.3)
        )
        pic1.line.color.rgb = BORDER_GRAY
        pic1.line.width = Pt(1)

    tb_p1 = s7.shapes.add_textbox(Inches(1.8), Inches(5.75), Inches(5.3), Inches(1.2))
    tf_p1 = tb_p1.text_frame
    tf_p1.word_wrap = True
    p = tf_p1.paragraphs[0]
    p.text = "Installation Pattern Designer"
    p.font.name = "Times New Roman"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RED_THEME
    p = tf_p1.add_paragraph()
    p.text = "Dynamic SVG visualizer for 6 layouts (Straight, Herringbone, Chevron, Basket Weave)."
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.color.rgb = BODY_TEXT

    # Screenshot 2: Calculator / Comparison
    if os.path.exists(shot_calc):
        pic2 = s7.shapes.add_picture(
            shot_calc,
            Inches(7.5), Inches(1.35), Inches(5.3), Inches(4.3)
        )
        pic2.line.color.rgb = BORDER_GRAY
        pic2.line.width = Pt(1)

    tb_p2 = s7.shapes.add_textbox(Inches(7.5), Inches(5.75), Inches(5.3), Inches(1.2))
    tf_p2 = tb_p2.text_frame
    tf_p2.word_wrap = True
    p = tf_p2.paragraphs[0]
    p.text = "Surface Area & Wastage Calculator"
    p.font.name = "Times New Roman"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RED_THEME
    p = tf_p2.add_paragraph()
    p.text = "Room dimension calculation, pattern-specific wastage formula (+5% to +18%), and cost estimation."
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.color.rgb = BODY_TEXT

    # =========================================================================
    # SLIDE 8: FUTURE SCOPE & KEY LEARNINGS
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s8)
    add_header_banner(s8, "Future Scope & Key Learnings")
    add_footer_banner(s8)

    # Left: Future Scope
    tb_fs = s8.shapes.add_textbox(Inches(1.8), Inches(1.3), Inches(5.3), Inches(5.6))
    tf_fs = tb_fs.text_frame
    tf_fs.word_wrap = True

    p = tf_fs.paragraphs[0]
    p.text = "Future Scope"
    p.font.name = "Times New Roman"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = RED_THEME

    future_points = [
        ("1. Headless CMS & Cloud API: ", "Integrate Strapi/Supabase for live catalog management and payment gateway."),
        ("2. WebGL 3D Room Visualizer: ", "Real-time 3D and Augmented Reality (AR) to project floorings directly into rooms."),
        ("3. AI Interior Recommendation: ", "Machine learning model analyzing room photos to suggest matching wood finishes."),
        ("4. B2B Trade Portal: ", "Contractor portal with tiered wholesale pricing and CAD/BIM downloads.")
    ]
    for num, txt in future_points:
        p = tf_fs.add_paragraph()
        p.font.name = "Calibri"
        p.font.size = Pt(15.5)
        p.line_spacing = 1.15
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = num
        r1.font.bold = True
        r2 = p.add_run()
        r2.text = txt

    # Right: Key Learnings
    tb_kl = s8.shapes.add_textbox(Inches(7.4), Inches(1.3), Inches(5.5), Inches(5.6))
    tf_kl = tb_kl.text_frame
    tf_kl.word_wrap = True

    p = tf_kl.paragraphs[0]
    p.text = "Key Industrial Learnings"
    p.font.name = "Times New Roman"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = RED_THEME

    learnings = [
        ("1. Component-Driven Architecture: ", "Structured reusable, decoupled UI components with strict TypeScript types."),
        ("2. Reactive Client-Side State: ", "Managed synchronized carts and sample orders using Context API and LocalStorage."),
        ("3. Interactive Tool Engineering: ", "Built dynamic SVG mathematical generators for architectural pattern simulation."),
        ("4. Production Standards: ", "Practiced Git workflows, code modularity, route code splitting, and sub-second load times.")
    ]
    for num, txt in learnings:
        p = tf_kl.add_paragraph()
        p.font.name = "Calibri"
        p.font.size = Pt(15.5)
        p.line_spacing = 1.15
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = num
        r1.font.bold = True
        r2 = p.add_run()
        r2.text = txt

    # =========================================================================
    # SLIDE 9: THANK YOU & Q&A
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_skit_logo(s9)
    add_footer_banner(s9)

    tb_end = s9.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.33), Inches(4.0))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p = tf_end.paragraphs[0]
    p.text = "THANK YOU"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = RED_THEME

    p = tf_end.add_paragraph()
    p.text = "Questions & Answers"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Times New Roman"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = BODY_TEXT
    p.space_before = Pt(16)

    p = tf_end.add_paragraph()
    p.text = "Lamiwood: Modern E-Commerce Platform for Architectural Surfaces"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(20)
    p.font.color.rgb = DARK_RED
    p.space_before = Pt(14)

    p = tf_end.add_paragraph()
    p.text = "Aditya Swami | 24ESKCS024 | SKIT Jaipur"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.color.rgb = BODY_TEXT
    p.space_before = Pt(6)

    # Save presentation
    prs.save(output_pptx)
    print(f"9-slide presentation saved to: {output_pptx}")

def export_to_pdf_and_images():
    ppt_app = win32com.client.Dispatch('PowerPoint.Application')
    presentation = ppt_app.Presentations.Open(os.path.abspath(output_pptx), WithWindow=False)

    # Clean previous images
    for f in os.listdir(slides_img_dir):
        if f.endswith('.png'):
            try: os.remove(os.path.join(slides_img_dir, f))
            except: pass

    # Export to PDF (32 is ppSaveAsPDF)
    presentation.SaveAs(os.path.abspath(output_pdf), 32)
    print(f"Exported PDF to: {output_pdf}")

    # Export slides to images
    for i, slide in enumerate(presentation.Slides):
        slide_img_path = os.path.join(slides_img_dir, f"slide_{i+1}.png")
        slide.Export(slide_img_path, "PNG", 1920, 1080)
        print(f"Exported slide {i+1} to {slide_img_path}")

    presentation.Close()
    ppt_app.Quit()
    print("All 9 slides exported successfully!")

if __name__ == "__main__":
    create_9_slide_presentation()
    export_to_pdf_and_images()
